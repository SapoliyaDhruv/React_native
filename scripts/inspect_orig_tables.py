import pptx

prs = pptx.Presentation('Review_Presentation_backup_orig.pptx')
for idx in [19, 20, 21, 22]:
    s = prs.slides[idx]
    print(f"\n=== Backup Slide {idx+1}: {s.shapes.title.text if s.shapes.title else ''} ===")
    for shp in s.shapes:
        if shp.has_table:
            t = shp.table
            print(f"Table pos: left={shp.left.inches:.2f}, top={shp.top.inches:.2f}, w={shp.width.inches:.2f}, h={shp.height.inches:.2f}")
            for r_idx in range(min(2, len(t.rows))):
                row_str = []
                for c_idx in range(len(t.columns)):
                    cell = t.cell(r_idx, c_idx)
                    p = cell.text_frame.paragraphs[0]
                    font_size = p.font.size.pt if p.font.size else "inherit"
                    font_color = p.font.color.rgb if p.font.color and p.font.color.type else "inherit"
                    row_str.append(f"{cell.text.strip()} (sz:{font_size}, col:{font_color})")
                print(f"  R{r_idx}: " + " | ".join(row_str))
            # Check cell XML for formatting / styling
            r0c0 = t.cell(0, 0)
            print("  R0C0 xml snippet:", r0c0._tc.xml[:200])
            r1c0 = t.cell(1, 0)
            print("  R1C0 xml snippet:", r1c0._tc.xml[:200])
        elif shp.has_text_frame:
            print("Text shape:", repr(shp.text_frame.text.strip()))

