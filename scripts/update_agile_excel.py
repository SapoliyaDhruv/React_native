import openpyxl

for fname in ['Agile Document (Recovered).xlsx', 'Agile Document_backup.xlsx']:
    try:
        wb = openpyxl.load_workbook(fname)
        count = 0
        for sname in wb.sheetnames:
            ws = wb[sname]
            for r in range(1, ws.max_row + 1):
                for c in range(1, ws.max_column + 1):
                    val = str(ws.cell(r, c).value or '')
                    if 'jignesh' in val.lower() or 'doshi' in val.lower() or 'jignash' in val.lower():
                        print(f'{fname} [{sname}] R{r}C{c}: was "{val}" -> set to empty')
                        ws.cell(r, c).value = ''
                        count += 1
        wb.save(fname)
        print(f'{fname}: modified {count} cells.')
    except Exception as e:
        print(f'Error updating {fname}: {e}')

# Verify no occurrences remain
for fname in ['Agile Document (Recovered).xlsx', 'Agile Document_backup.xlsx']:
    wb = openpyxl.load_workbook(fname, data_only=True)
    found = []
    for sname in wb.sheetnames:
        ws = wb[sname]
        for r in range(1, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                val = str(ws.cell(r, c).value or '')
                if 'jignesh' in val.lower() or 'doshi' in val.lower() or 'jignash' in val.lower():
                    found.append((sname, r, c, val))
    print(f'{fname} verification: {len(found)} occurrences remaining.')

