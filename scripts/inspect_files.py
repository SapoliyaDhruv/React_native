import pptx
import openpyxl
import docx

def inspect_pptx(filename):
    print(f"=== {filename} ===")
    prs = pptx.Presentation(filename)
    print(f"Total slides: {len(prs.slides)}")
    for idx, slide in enumerate(prs.slides):
        title = ""
        if slide.shapes.title and slide.shapes.title.text:
            title = slide.shapes.title.text.strip().replace('\n', ' ')
        shapes_info = []
        for s in slide.shapes:
            if s.has_text_frame:
                txt = s.text_frame.text.strip().replace('\n', ' ')
                if txt:
                    shapes_info.append(f"Text({txt[:50]}...)")
            elif s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
                shapes_info.append("Picture")
            elif s.has_table:
                shapes_info.append("Table")
            else:
                shapes_info.append(f"Shape({s.name})")
        print(f"Slide {idx+1}: '{title}' | {shapes_info}")

def inspect_xlsx(filename):
    print(f"\n=== {filename} ===")
    wb = openpyxl.load_workbook(filename, data_only=True)
    print("Sheets:", wb.sheetnames)
    for name in wb.sheetnames:
        ws = wb[name]
        print(f"--- Sheet: {name} ({ws.max_row} rows, {ws.max_column} cols) ---")
        for r in range(1, min(ws.max_row+1, 10)):
            vals = [str(ws.cell(r, c).value or '')[:20] for c in range(1, min(ws.max_column+1, 8))]
            if any(vals):
                print(f"  R{r}: {' | '.join(vals)}")

if __name__ == "__main__":
    inspect_pptx("Review_Presentation.pptx")
    inspect_xlsx("Agile Document (Recovered).xlsx")

