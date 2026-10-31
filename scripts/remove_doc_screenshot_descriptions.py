import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import shutil

def update_screenshots_in_doc():
    doc_path = "MindSprint_Project_Documentation.docx"
    # Backup first
    shutil.copy(doc_path, "MindSprint_Project_Documentation_backup_before_desc_removal.docx")
    
    doc = docx.Document(doc_path)
    
    # Locate "4.2 Screen Shots" and "5. Agile Documentation"
    idx_start = -1
    idx_end = -1
    
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if t == "4.2 Screen Shots":
            idx_start = i
        elif t == "5. Agile Documentation":
            idx_end = i
            break
            
    print(f"Found '4.2 Screen Shots' at paragraph {idx_start} and '5. Agile Documentation' at {idx_end}")
    assert idx_start != -1 and idx_end != -1, "Could not find start/end paragraphs"
    
    p_agile = doc.paragraphs[idx_end]
    
    # Collect all paragraph elements strictly between idx_start and idx_end to remove
    to_remove = [doc.paragraphs[k] for k in range(idx_start + 1, idx_end)]
    print(f"Removing {len(to_remove)} old paragraphs between 4.2 and 5.0...")
    for p in to_remove:
        sp = p._element
        sp.getparent().remove(sp)
        
    # Now define the 10 clean screenshot items (NO descriptions!)
    screenshots_data = [
        ("4.2.1 Welcome and Onboarding Screen",
         "screenshots/01_onboarding.png",
         "Figure 4.1: Welcome & Onboarding Screen - MindSprint Academic & Career Companion"),
        ("4.2.2 User Authentication and Sign-in Screen",
         "screenshots/02_login.png",
         "Figure 4.2: User Authentication & Sign-In Interface"),
        ("4.2.3 Student Dashboard & Learning Hub",
         "screenshots/03_dashboard_home.png",
         "Figure 4.3: Main Dashboard, Syllabus Selector & Gamified Study Hub"),
        ("4.2.4 Course Catalog & Chapter Navigator",
         "screenshots/04_course_catalog.png",
         "Figure 4.4: Course Catalog, Chapter Navigator & Official e-Textbook Reader"),
        ("4.2.5 Interactive Video Lecture Player",
         "screenshots/05_video_learning.png",
         "Figure 4.5: High-Definition Video Masterclasses & Lecture Player"),
        ("4.2.6 Real-Time Timed Mock Exam Engine",
         "screenshots/06_mock_exam.png",
         "Figure 4.6: Real-Time Mock Exam Simulation & Question Palette"),
        ("4.2.7 Spaced Repetition Flashcards",
         "screenshots/08_flashcards.png",
         "Figure 4.7: Interactive Flashcards & Spaced Repetition Recall"),
        ("4.2.8 AI Doubt Solver & Camera Scanner",
         "screenshots/10_doubt_solver.png",
         "Figure 4.8: AI-Powered Doubt Solver & OCR Question Matching"),
        ("4.2.9 Campus Leaderboard & Real-Time Rankings",
         "screenshots/11_leaderboard.png",
         "Figure 4.9: Campus Leaderboard, Study XP & Competitive Peer Rankings"),
        ("4.2.10 Student Profile, Badges & Theme Configuration",
         "screenshots/13_student_profile.png",
         "Figure 4.10: Student Profile, Gamification Badges & Theme Configuration")
    ]
    
    # Insert each item before p_agile
    for title_text, img_path, caption_text in screenshots_data:
        # 1. Subsection title
        p_title = p_agile.insert_paragraph_before()
        p_title.paragraph_format.space_before = Pt(14)
        p_title.paragraph_format.space_after = Pt(4)
        p_title.paragraph_format.keep_with_next = True
        run_title = p_title.add_run(title_text)
        run_title.bold = True
        run_title.font.name = "Arial"
        run_title.font.size = Pt(12)
        run_title.font.color.rgb = RGBColor(15, 23, 42)
        
        # 2. Image
        p_img = p_agile.insert_paragraph_before()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(4)
        p_img.paragraph_format.space_after = Pt(3)
        p_img.paragraph_format.keep_with_next = True
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(2.8))
        
        # 3. Caption
        p_cap = p_agile.insert_paragraph_before()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(12)
        run_cap = p_cap.add_run(caption_text)
        run_cap.italic = True
        run_cap.font.name = "Calibri"
        run_cap.font.size = Pt(9.5)
        run_cap.font.color.rgb = RGBColor(100, 116, 139)
        
    doc.save(doc_path)
    print(f"Successfully updated {doc_path} without screenshot descriptions! Total paragraphs: {len(doc.paragraphs)}")

if __name__ == "__main__":
    update_screenshots_in_doc()

