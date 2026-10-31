import docx

doc = docx.Document('ProjectReviewSheet-ReactNative.docx')
t = doc.tables[0]

t.cell(2, 0).text = "Project Name: MindSprint - Placement Preparation & Skill Learning App"
t.cell(8, 3).text = "Sapoliya Dhruv"
t.cell(9, 3).text = "Vavaliya Dvarkesh"
t.cell(10, 3).text = "Prajapati Krunal"

doc.save('ProjectReviewSheet-ReactNative.docx')
print("Updated ProjectReviewSheet-ReactNative.docx successfully!")

