import docx

doc = docx.Document('MindSprint_Project_Documentation.docx')
print(f"Total paragraphs: {len(doc.paragraphs)}, Total tables: {len(doc.tables)}")
for i in range(165, len(doc.paragraphs)):
    p = doc.paragraphs[i]
    r_imgs = [r for r in p.runs if 'graphic' in r._element.xml]
    img_tag = f" [HAS_IMAGE: {len(r_imgs)}]" if r_imgs else ""
    print(f"P{i:3d}{img_tag}: {p.text[:75]}")

