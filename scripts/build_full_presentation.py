import pptx
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
import os

def build_presentation():
    input_file = "Review_Presentation_backup_orig.pptx"
    output_file = "Review_Presentation.pptx"

    prs = pptx.Presentation(input_file)
    print(f"Loaded {input_file} with {len(prs.slides)} slides")

    # 1. Update text on existing slides
    # Fix Slide 5 (orig idx 4)
    s5 = prs.slides[4]
    for shp in s5.shapes:
        if shp.has_text_frame and "5To offer" in shp.text_frame.text:
            shp.text_frame.text = shp.text_frame.text.replace("5To offer", "To offer")
            print("Fixed typo on Slide 5")

    # Update Slide 7 (orig idx 6) - Core Components
    s7 = prs.slides[6]
    for shp in s7.shapes:
        if shp.has_text_frame and shp != s7.shapes.title:
            tf = shp.text_frame
            tf.word_wrap = True
            tf.clear()
            p0 = tf.paragraphs[0]
            p0.text = "MindSprint is architected into five core software layers and modules:"
            p0.font.bold = True
            p0.font.size = Pt(13)

            components = [
                ("1. Mobile Client App:", " React Native frontend built with Expo v54/v57, TypeScript, and Expo Router, delivering cross-platform performance across Android, iOS, and Web."),
                ("2. REST API Backend:", " Node.js and Express server delivering secure JWT authentication, rate limiting, academic data APIs, and AI proxy routing."),
                ("3. Scalable Database Layer:", " MongoDB database with Mongoose schemas for User accounts, Question banks, Courses, ExamResults, DuelRooms, and SamplePapers."),
                ("4. Offline Resilience Layer:", " Service-level fallback mechanism (ApiService) that automatically loads verified local datasets when network connectivity drops."),
                ("5. Interactive Learning Engines:", " Real-time timed mock exam engine with question palette, 1v1 Quiz Duel battles, Voice Viva player, and AI Doubt Solver with photo OCR.")
            ]
            for title, desc in components:
                p = tf.add_paragraph()
                p.space_before = Pt(8)
                run1 = p.add_run()
                run1.text = title
                run1.font.bold = True
                run1.font.size = Pt(11.5)
                run2 = p.add_run()
                run2.text = desc
                run2.font.size = Pt(11)
            print("Updated Slide 7 with core components")

    # Update Slide 12 (orig idx 11) - Functional Requirements
    s12 = prs.slides[11]
    for shp in s12.shapes:
        if shp.has_text_frame and shp != s12.shapes.title:
            tf = shp.text_frame
            tf.word_wrap = True
            tf.clear()
            p0 = tf.paragraphs[0]
            p0.text = "Functional Requirements :"
            p0.font.bold = True
            p0.font.size = Pt(13)

            reqs = [
                ("1. User Authentication:", " Secure registration, login, and guest mode protected by bcrypt password hashing and JWT tokens."),
                ("2. Multi-Board Track Selection:", " Dynamic curriculum switching between GTU Engineering and Gujarat School (GSEB) syllabi."),
                ("3. Company-Wise Question Bank:", " Practice questions filtered by company recruitment pattern (TCS, Infosys, Amazon, etc.)."),
                ("4. Timed Mock Exam Engine:", " Real-time countdown timer, color-coded question palette navigation, and in-exam scratchpad."),
                ("5. Automated Weak Spot Detection:", " Automated score calculation flagging topics with <60% accuracy for targeted revision."),
                ("6. AI Doubt Solver & OCR:", " Instant question clearing via camera photo OCR and Gemini-powered step-by-step verified explanations."),
                ("7. Course Catalog & Video Masterclasses:", " Organized chapter navigator, official e-textbooks, and curated YouTube video lectures."),
                ("8. Engagement & 1v1 PvP Battles:", " Swipeable flashcards, daily study streaks, campus leaderboard, and live Quiz Duel matchmaking.")
            ]
            for title, desc in reqs:
                p = tf.add_paragraph()
                p.space_before = Pt(6)
                run1 = p.add_run()
                run1.text = title
                run1.font.bold = True
                run1.font.size = Pt(11)
                run2 = p.add_run()
                run2.text = desc
                run2.font.size = Pt(10.5)
            print("Updated Slide 12 with functional requirements")

    # Helper function to clear non-title shapes
    def clear_non_title_shapes(slide):
        to_del = [shp for shp in slide.shapes if shp != slide.shapes.title]
        for shp in to_del:
            sp = shp._element
            sp.getparent().remove(sp)

    # 2. Update 4.2 Screenshots on orig indices 26..30 (5 slides, 10 screenshots total, 2 per slide, NO text info)
    screenshot_slides_data = [
        ("4.2 Screen Shots: Onboarding & Authentication",
         "screenshots/01_onboarding.png", "screenshots/02_login.png"),
        ("4.2 Screen Shots: Student Dashboard & Course Catalog",
         "screenshots/03_dashboard_home.png", "screenshots/04_course_catalog.png"),
        ("4.2 Screen Shots: Video Learning & Timed Mock Exam",
         "screenshots/05_video_learning.png", "screenshots/06_mock_exam.png"),
        ("4.2 Screen Shots: Spaced Repetition Flashcards & AI Doubt Solver",
         "screenshots/08_flashcards.png", "screenshots/10_doubt_solver.png"),
        ("4.2 Screen Shots: Campus Leaderboard & Student Profile",
         "screenshots/11_leaderboard.png", "screenshots/13_student_profile.png")
    ]
    for i, (title_text, img1, img2) in enumerate(screenshot_slides_data):
        slide = prs.slides[26 + i]
        clear_non_title_shapes(slide)
        if slide.shapes.title:
            slide.shapes.title.text = title_text
        slide.shapes.add_picture(img1, Inches(1.8), Inches(1.05), Inches(2.65), Inches(4.3))
        slide.shapes.add_picture(img2, Inches(5.55), Inches(1.05), Inches(2.65), Inches(4.3))
        print(f"Updated Screenshot Slide {i+1}: {title_text}")

    # 3. Update Agile Slides on orig indices 35..42 (8 slides: 5.1 to 5.8)
    agile_slides_data = [
        ("5.1 Agile Project Charter", "screenshots/agile_01_project_charter.png"),
        ("5.2 Agile Roadmap / Schedule", "screenshots/agile_02_project_roadmap.png"),
        ("5.3 Agile Project Plan", "screenshots/agile_03_project_plan.png"),
        ("5.4 Agile User Stories (Minimum 3 Tasks)", "screenshots/agile_04_user_stories.png"),
        ("5.5 Agile Release Plan", "screenshots/agile_05_release_plan.png"),
        ("5.6 Agile Product Backlog & Sprint Backlog", "screenshots/agile_06_product_backlog.png"),
        ("5.7 Agile Test Plan", "screenshots/agile_07_test_plan.png"),
        ("5.8 Earned-Value and Burn Charts", "screenshots/agile_08_burndown_chart.png")
    ]
    for i, (title_text, img_path) in enumerate(agile_slides_data):
        slide = prs.slides[35 + i]
        clear_non_title_shapes(slide)
        if slide.shapes.title:
            slide.shapes.title.text = title_text
        slide.shapes.add_picture(img_path, Inches(0.6), Inches(1.05), Inches(8.8), Inches(4.35))
        print(f"Updated Agile Slide {i+1}: {title_text}")

    # Helper function to populate content slide
    def populate_content_slide(slide, title_text, paragraphs_data):
        clear_non_title_shapes(slide)
        if slide.shapes.title:
            slide.shapes.title.text = title_text
            slide.shapes.title.left = Inches(0.65)
            slide.shapes.title.top = Inches(0.4)
            slide.shapes.title.width = Inches(8.7)
            slide.shapes.title.height = Inches(0.6)

        tx_box = slide.shapes.add_textbox(Inches(0.65), Inches(1.15), Inches(8.7), Inches(4.2))
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        first = True
        for heading, body in paragraphs_data:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            if heading:
                p.space_before = Pt(8)
                r_head = p.add_run()
                r_head.text = heading + " "
                r_head.font.bold = True
                r_head.font.size = Pt(11.5)
                r_head.font.color.rgb = RGBColor(30, 58, 138)
            if body:
                r_body = p.add_run()
                r_body.text = body
                r_body.font.size = Pt(11)
                r_body.font.color.rgb = RGBColor(51, 65, 85)

    # 4. Populate Content Slides on orig indices 31, 32, 33 (which were old screenshot slides, layout TITLE_ONLY)
    # Orig idx 31 -> Section 6 Content: Proposed Enhancements
    enhancements_content = [
        ("1. Adaptive AI Difficulty Engine:", "Dynamically adjusts mock exam question difficulty in real time based on individual student performance history, accuracy streaks, and time spent per question."),
        ("2. Real-Time Multiplayer Quiz Battles:", "Expands the 1v1 Quiz Duel into live classroom multiplayer tournaments and collaborative study groups with live voice viva discussion rooms."),
        ("3. Placement Cell & Faculty Analytics Portal:", "Web-based administration portal enabling college placement officers to monitor cohort preparedness, track weak subject patterns across student batches, and schedule targeted mock drives."),
        ("4. Background Exam Sync & Smart Push Alerts:", "Automatic offline-to-online test score syncing upon reconnection, paired with intelligent push notifications to maintain daily study streak retention."),
        ("5. Regional Vernacular Language Support:", "Adding bilingual Gujarati and Hindi audio and explanations for academic school syllabi and technical definitions alongside standard English.")
    ]
    populate_content_slide(prs.slides[31], "6. Proposed Enhancements", enhancements_content)
    print("Populated orig idx 31 with Section 6 Content")

    # Orig idx 32 -> Section 7 Content: Conclusion
    conclusion_content = [
        ("• Unified Learning Solution:", "MindSprint successfully eliminates the fragmentation of student preparation by integrating aptitude mock tests, technical video masterclasses, formula flashcards, and AI doubt resolution into a single responsive mobile application."),
        ("• Production-Grade Full-Stack Architecture:", "Built with React Native, Expo, TypeScript, Node.js REST APIs, and MongoDB, supported by an intelligent offline fallback layer that guarantees zero downtime and uninterrupted learning during network drops."),
        ("• Proven Agile Execution:", "Developed iteratively across six sprints using the Scrum framework, incorporating sprint backlogs, burn-down velocity tracking, and comprehensive quality assurance testing across Android and Web environments."),
        ("• Direct Academic & Placement Impact:", "Automated weak-spot diagnostics (<60% accuracy flagging) and bite-sized flashcards empower college students to pinpoint critical improvement areas and excel in campus placement drives.")
    ]
    populate_content_slide(prs.slides[32], "7. Conclusion", conclusion_content)
    print("Populated orig idx 32 with Section 7 Content")

    # Orig idx 33 -> Section 8 Content: Bibliography
    bibliography_content = [
        ("1. Expo & React Native Framework Documentation:", "Official architecture guides, Expo SDK v54/v57 APIs, Expo Router file-based routing, and React Native component lifecycle. https://docs.expo.dev"),
        ("2. Node.js & Express.js Developer Guides:", "RESTful API design principles, Express middleware pipelines, JWT bearer token authentication, and bcryptjs password security. https://expressjs.com"),
        ("3. MongoDB & Mongoose Object Data Modeling:", "NoSQL document schema design, indexing strategies, and MongoDB Atlas database connection pooling. https://mongoosejs.com"),
        ("4. Google Gemini API Documentation:", "Multimodal generative AI integration, prompt engineering, and image OCR doubt solving. https://ai.google.dev"),
        ("5. Gujarat Technological University (GTU) & GSEB Curriculum Manuals:", "Official engineering syllabus structures, academic exam blueprints, and subject textbooks."),
        ("6. The Scrum Guide (Ken Schwaber & Jeff Sutherland):", "The definitive rules and guidelines of the Scrum framework for Agile software development and sprint planning.")
    ]
    populate_content_slide(prs.slides[33], "8. Bibliography", bibliography_content)
    print("Populated orig idx 33 with Section 8 Content")

    # 5. Update Thank You Slide on orig idx 43
    s_thank = prs.slides[43]
    clear_non_title_shapes(s_thank)
    if s_thank.shapes.title:
        s_thank.shapes.title.text = "Thank You!"
    tb = s_thank.shapes.add_textbox(Inches(1.0), Inches(2.6), Inches(8.0), Inches(2.2))
    tf = tb.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = "Questions & Discussion"
    p1.alignment = PP_ALIGN.CENTER
    p1.font.size = Pt(20)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(30, 58, 138)

    p2 = tf.add_paragraph()
    p2.space_before = Pt(14)
    p2.text = "Project Team Members:\nSapoliya Dhruv  |  Vavaliya Dvarkesh  |  Prajapati Krunal"
    p2.alignment = PP_ALIGN.CENTER
    p2.font.size = Pt(14)
    p2.font.color.rgb = RGBColor(51, 65, 85)

    p3 = tf.add_paragraph()
    p3.space_before = Pt(10)
    p3.text = "Lok Jagruti Institute of Computer Applications (LJICA)"
    p3.alignment = PP_ALIGN.CENTER
    p3.font.size = Pt(12)
    p3.font.color.rgb = RGBColor(100, 116, 139)
    print("Updated Thank You slide on orig idx 43")

    # 6. Add the 3 Section Header slides (layout 4: BIG_NUMBER)
    s_h6 = prs.slides.add_slide(prs.slide_layouts[4])
    s_h6.shapes.title.text = "6. Proposed Enhancements"

    s_h7 = prs.slides.add_slide(prs.slide_layouts[4])
    s_h7.shapes.title.text = "7. Conclusion"

    s_h8 = prs.slides.add_slide(prs.slide_layouts[4])
    s_h8.shapes.title.text = "8. Bibliography"
    print("Added 3 section header slides")

    # 7. Now reorder slides in _sldIdLst to form the exact, beautiful presentation
    # Elements in sldIdLst right now (total 47):
    # orig 0..43 (indices 0..43)
    # new_1 (idx 44, Header 6)
    # new_2 (idx 45, Header 7)
    # new_3 (idx 46, Header 8)
    sldIdLst = prs.slides._sldIdLst
    all_elems = list(sldIdLst)

    # Desired order:
    # 0..25: orig 0..25 (Title, Intro, Requirements, Design, Coding standards)
    # 26..30: orig 26..30 (Screenshots 1..5)
    # 31: orig 34 (Header: 5. Agile Documentation)
    # 32..39: orig 35..42 (Agile 5.1 to 5.8)
    # 40: new_1 (all_elems[44], Header: 6. Proposed Enhancements)
    # 41: orig 31 (all_elems[31], Content: 6. Proposed Enhancements)
    # 42: new_2 (all_elems[45], Header: 7. Conclusion)
    # 43: orig 32 (all_elems[32], Content: 7. Conclusion)
    # 44: new_3 (all_elems[46], Header: 8. Bibliography)
    # 45: orig 33 (all_elems[33], Content: 8. Bibliography)
    # 46: orig 43 (all_elems[43], Thank You!)

    new_order = (
        all_elems[0:26] +        # 26 slides (Title to Coding Standards)
        all_elems[26:31] +       # 5 screenshot slides
        [all_elems[34]] +        # Header 5: Agile Documentation
        all_elems[35:43] +       # 8 Agile slides
        [all_elems[44]] +        # Header 6: Proposed Enhancements
        [all_elems[31]] +        # Content 6: Proposed Enhancements
        [all_elems[45]] +        # Header 7: Conclusion
        [all_elems[32]] +        # Content 7: Conclusion
        [all_elems[46]] +        # Header 8: Bibliography
        [all_elems[33]] +        # Content 8: Bibliography
        [all_elems[43]]          # Thank You!
    )

    sldIdLst.clear()
    for elem in new_order:
        sldIdLst.append(elem)

    prs.save(output_file)
    print(f"Presentation successfully saved to {output_file}! Total slides: {len(prs.slides)}")

if __name__ == "__main__":
    build_presentation()

