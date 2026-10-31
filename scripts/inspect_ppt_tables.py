import pptx

prs = pptx.Presentation('Review_Presentation.pptx')
for idx in [19, 20, 21, 22]:
    s = prs.slides[idx]
    print(f"\n=== Slide {idx+1}: {s.shapes.title.text if s.shapes.title else 'No Title'} ===")
    for shp in s.shapes:
        if shp.has_table:
            t = shp.table
            print(f"Table: {len(t.rows)} rows, {len(t.columns)} cols, left={shp.left.inches}, top={shp.top.inches}, w={shp.width.inches}, h={shp.height.inches}")
            for r_idx, r in enumerate(t.rows):
                txts = [c.text.strip().replace('\n', ' ') for c in r.cells]
                print(f"  R{r_idx}: " + " | ".join(txts))
        elif shp.has_text_frame and shp != s.shapes.title:
            print("  Text box:", shp.text_frame.text[:60])

