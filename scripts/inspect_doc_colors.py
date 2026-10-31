import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

doc = docx.Document('MindSprint_Project_Documentation.docx')
colored_runs = []
for p_idx, p in enumerate(doc.paragraphs):
    for r_idx, r in enumerate(p.runs):
        if r.font.color and r.font.color.rgb:
            colored_runs.append((p_idx, r.text[:30], str(r.font.color.rgb)))

print(f"Total colored text runs in paragraphs: {len(colored_runs)}")
for item in colored_runs[:15]:
    print(f"  P{item[0]}: '{item[1]}' -> {item[2]}")

# Check tables
colored_cells = 0
for t_idx, t in enumerate(doc.tables):
    for r_idx, row in enumerate(t.rows):
        for c_idx, cell in enumerate(row.cells):
            tcPr = cell._element.tcPr
            if tcPr is not None:
                shd = tcPr.find(qn('w:shd'))
                if shd is not None and shd.get(qn('w:fill')) not in [None, 'auto', 'FFFFFF', 'ffffff']:
                    colored_cells += 1

print(f"Total shaded cells in tables: {colored_cells}")

