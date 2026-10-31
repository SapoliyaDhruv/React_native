import docx
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import parse_xml
from docx.oxml.ns import qn
import shutil

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    # Remove existing w:shd if any
    for shd in tcPr.findall(qn('w:shd')):
        tcPr.remove(shd)
    shd = parse_xml(f'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:val="clear" w:color="auto" w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def apply_bw_borders(tbl):
    tblPr = tbl._element.xpath('w:tblPr')
    if tblPr:
        for b in tblPr[0].xpath('w:tblBorders'):
            tblPr[0].remove(b)
        borders = parse_xml(
            '<w:tblBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            '<w:top w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
            '<w:left w:val="single" w:sz="6" w:space="0" w:color="666666"/>'
            '<w:bottom w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
            '<w:right w:val="single" w:sz="6" w:space="0" w:color="666666"/>'
            '<w:insideH w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
            '<w:insideV w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
            '</w:tblBorders>'
        )
        tblPr[0].append(borders)

def make_doc_black_and_white():
    doc_path = "MindSprint_Project_Documentation.docx"
    shutil.copy(doc_path, "MindSprint_Project_Documentation_backup_before_bw.docx")
    doc = docx.Document(doc_path)
    
    print(f"Processing {doc_path} ({len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables)...")
    
    # 1. Update all paragraphs to pure black text
    for p in doc.paragraphs:
        for r in p.runs:
            r.font.color.rgb = RGBColor(0, 0, 0)
            
    # 2. Update all tables
    for t_idx, tbl in enumerate(doc.tables):
        apply_bw_borders(tbl)
        for r_idx, row in enumerate(tbl.rows):
            is_header = (r_idx == 0)
            fill_color = "E5E7EB" if is_header else "FFFFFF"
            
            for c_idx, cell in enumerate(row.cells):
                set_cell_background(cell, fill_color)
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.color.rgb = RGBColor(0, 0, 0)
                        if is_header:
                            r.bold = True
                            
    doc.save(doc_path)
    print(f"Successfully converted {doc_path} to complete Black & White!")

if __name__ == "__main__":
    make_doc_black_and_white()

