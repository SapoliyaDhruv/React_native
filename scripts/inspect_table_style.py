import docx

doc = docx.Document('MindSprint_Project_Documentation.docx')
t = doc.tables[0]
print("Table style:", t.style.name if t.style else "No style")
print("Table alignment:", t.alignment)
for c_idx in range(len(t.columns)):
    cell = t.cell(0, c_idx)
    print(f"Col {c_idx} width:", cell.width.inches if cell.width else "auto", "text:", cell.text)

