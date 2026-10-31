import pptx

prs = pptx.Presentation('Review_Presentation.pptx')
print(f"Total slides: {len(prs.slides)}")
for idx, s in enumerate(prs.slides):
    title = s.shapes.title.text.strip().replace('\n', ' ') if s.shapes.title else 'NO TITLE'
    shape_types = [shp.shape_type.name for shp in s.shapes]
    print(f"Slide {idx+1:2d}: [{s.slide_layout.name}] '{title}' -> {shape_types}")

