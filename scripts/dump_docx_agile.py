import docx

doc = docx.Document('MindSprint_Project_Documentation.docx')
for i in range(169, len(doc.paragraphs)):
    p = doc.paragraphs[i]
    if p.text.strip():
        print(f"P{i}: {p.text}")

