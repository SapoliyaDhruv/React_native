import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import shutil

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def populate_table(tbl, rows_data):
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    col_widths = [Inches(1.5), Inches(1.3), Inches(1.6), Inches(2.1)]
    
    for r_idx, row_values in enumerate(rows_data):
        row = tbl.rows[r_idx]
        for c_idx, val in enumerate(row_values):
            cell = row.cells[c_idx]
            cell.width = col_widths[c_idx]
            cell.text = val
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.05
            
            run = p.runs[0]
            if r_idx == 0:
                set_cell_background(cell, '1E3A8A') # Navy blue header
                set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
                run.bold = True
                run.font.name = 'Arial'
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(255, 255, 255)
            else:
                fill_color = 'FFFFFF' if r_idx % 2 == 1 else 'F8FAFC'
                set_cell_background(cell, fill_color)
                set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
                run.font.name = 'Calibri'
                run.font.size = Pt(9.0)
                if c_idx == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor(30, 58, 138)
                elif c_idx == 1:
                    run.font.color.rgb = RGBColor(14, 116, 144)
                elif c_idx == 2:
                    run.font.color.rgb = RGBColor(180, 83, 9)
                else:
                    run.font.color.rgb = RGBColor(51, 65, 85)

def update_data_dictionary():
    doc_path = "MindSprint_Project_Documentation.docx"
    shutil.copy(doc_path, "MindSprint_Project_Documentation_backup_before_dd.docx")
    doc = docx.Document(doc_path)

    # 11 full collection schemas
    all_schemas = [
        ("Table 1: User Schema (users)", [
            ["Field Name", "Data Type", "Constraint", "Description"],
            ["_id", "ObjectId", "Primary Key", "Unique user identifier"],
            ["name", "String", "Required", "Full name of the student"],
            ["email", "String", "Unique, Required", "Student email address"],
            ["passwordHash", "String", "Required", "Bcrypt-hashed password string"],
            ["activeTrack", "String", "Default: 'placement'", "Selected track ('placement', 'school')"],
            ["grade", "String", "Enum: 6-12, 'college'", "Current academic grade / year"],
            ["board", "String", "Enum: CBSE, GSEB, ICSE", "Academic education board"],
            ["medium", "String", "Enum: English, Gujarati", "Medium of instruction"],
            ["streakCount", "Number", "Default: 3", "Daily consecutive study streak"],
            ["targetCompany", "String", "Default: 'All Companies'", "Target placement employer"],
            ["createdAt", "Date", "Default: Now", "Account creation timestamp"]
        ]),
        ("Table 2: Question Schema (questions)", [
            ["Field Name", "Data Type", "Constraint", "Description"],
            ["_id", "ObjectId", "Primary Key", "Unique question identifier"],
            ["id", "Number", "Unique, Indexed", "Numeric question index"],
            ["questionText", "String", "Required", "The practice question statement"],
            ["options", "Array[String]", "Required, Len: 4", "Multiple-choice answer options"],
            ["correctAnswer", "Number", "Required", "0-based index of correct option"],
            ["category", "String", "Required, Indexed", "Aptitude, Logical, Technical topic"],
            ["company", "String", "Optional, Indexed", "Target company tag (TCS, Infosys)"],
            ["difficulty", "String", "Enum: Easy, Med, Hard", "Calibrated difficulty level"],
            ["explanation", "String", "Optional", "Detailed step-by-step solution"],
            ["grades", "Array[String]", "Indexed", "Target grades (['college'], ['10'])"]
        ]),
        ("Table 3: Course Schema (courses)", [
            ["Field Name", "Data Type", "Constraint", "Description"],
            ["_id", "ObjectId", "Primary Key", "Unique course identifier"],
            ["title", "String", "Required", "Title of the IT course / subject"],
            ["category", "String", "Required", "Course domain (DSA, Web Dev, Cloud)"],
            ["instructor", "String", "Required", "Course instructor / educator name"],
            ["duration", "String", "Default: '10 hrs'", "Total instructional video hours"],
            ["lessonsCount", "Number", "Default: 20", "Total number of video lessons"],
            ["rating", "Number", "Default: 4.8", "Student rating out of 5.0"],
            ["tag", "String", "Default: 'Trending'", "Display badge tag"],
            ["level", "String", "Default: 'All Levels'", "Target student skill level"],
            ["thumbnailUrl", "String", "Optional", "Course cover graphic image URL"]
        ]),
        ("Table 4: ExamResult Schema (examresults)", [
            ["Field Name", "Data Type", "Constraint", "Description"],
            ["_id", "ObjectId", "Primary Key", "Unique result identifier"],
            ["userId", "ObjectId", "Foreign Key (User)", "Reference to student account"],
            ["userEmail", "String", "Default: Guest", "Student email address"],
            ["grade", "String", "Default: 'college'", "Academic grade level"],
            ["board", "String", "Default: 'CBSE'", "Examination board context"],
            ["category", "String", "Default: 'General'", "Subject or company test category"],
            ["scorePercentage", "Number", "Required", "Calculated score percentage (0-100)"],
            ["correctCount", "Number", "Required", "Number of correctly answered questions"],
            ["totalQuestions", "Number", "Required", "Total test questions count"],
            ["userAnswers", "Object", "Default: {}", "Key-value pairs of student answers"],
            ["weakCategories", "Array[String]", "Optional", "Topics where accuracy fell <60%"],
            ["createdAt", "Date", "Default: Now", "Exam submission timestamp"]
        ]),
        ("Table 5: DuelRoom Schema (duelrooms)", [
            ["Field Name", "Data Type", "Constraint", "Description"],
            ["_id", "ObjectId", "Primary Key", "Unique duel room identifier"],
            ["roomCode", "String", "Unique, Indexed", "6-character room code for 1v1 PvP"],
            ["grade", "String", "Required", "Academic grade level"],
            ["subject", "String", "Required, Default: 'All'", "Test subject filter"],
            ["status", "String", "Enum, Indexed", "waiting, ready, active, completed, abandoned"],
            ["host", "Object", "Required", "Host details (name, avatar, score, ready)"],
            ["guest", "Object", "Default: null", "Opponent student details"],
            ["questions", "Array[Object]", "Required", "Synchronized test questions for duel"],
            ["createdAt", "Date", "TTL: 3600s (1 hr)", "Auto-expiring session timestamp"]
        ]),
        ("Table 6: SamplePaper Schema (samplepapers)", [
            ["Field Name", "Data Type", "Constraint", "Description"],
            ["_id", "ObjectId", "Primary Key", "Unique sample paper identifier"],
            ["id", "String", "Unique, Required", "Model paper identifier code"],
            ["title", "String", "Required", "Official model paper title"],
            ["grade", "String", "Required, Indexed", "Target grade ('10', '12', 'college')"],
            ["board", "String", "Default: 'CBSE', Indexed", "Board (CBSE, GSEB, Placement)"],
            ["subject", "String", "Required, Indexed", "Examination subject"],
            ["year", "String", "Default: '2024-2025'", "Academic examination year"],
            ["durationMinutes", "Number", "Default: 60", "Total allocated exam duration"],
            ["totalMarks", "Number", "Default: 50", "Total examination marks"],
            ["questionCount", "Number", "Default: 25", "Total question count"],
            ["officialPdfUrl", "String", "Optional", "URL to downloadable paper PDF"]
        ]),
        ("Table 7: VideoLecture Schema (videolectures)", [
            ["Field Name", "Data Type", "Constraint", "Description"],
            ["_id", "ObjectId", "Primary Key", "Unique lecture identifier"],
            ["id", "String", "Unique, Required", "Video lesson identifier code"],
            ["grades", "Array[String]", "Indexed", "Grade tags (['college'], ['10'])"],
            ["boards", "Array[String]", "Indexed", "Educational boards (CBSE, GSEB)"],
            ["category", "String", "Required, Indexed", "Subject or technical topic category"],
            ["title", "String", "Required", "Lecture lesson title"],
            ["instructor", "String", "Required", "Course educator name"],
            ["duration", "String", "Required", "Formatted length (e.g. '18m 30s')"],
            ["durationSeconds", "Number", "Required", "Total playback duration in seconds"],
            ["youtubeVideoId", "String", "Required", "Embedded YouTube video ID"],
            ["thumbnailUrl", "String", "Required", "Thumbnail image URL"]
        ]),
        ("Table 8: Flashcard Schema (flashcards)", [
            ["Field Name", "Data Type", "Constraint", "Description"],
            ["_id", "ObjectId", "Primary Key", "Unique flashcard identifier"],
            ["id", "String", "Unique, Required", "Concept flashcard code"],
            ["topic", "String", "Required", "Subject topic (SQL, OS, Aptitude)"],
            ["question", "String", "Required", "Concept prompt or interview question"],
            ["answer", "String", "Required", "Concise verified answer / definition"],
            ["interviewerTips", "String", "Optional", "Practical advice for campus interview"],
            ["difficulty", "String", "Enum: Beginner, Med, Adv", "Topic complexity level"],
            ["tags", "Array[String]", "Optional", "Search keywords and category tags"],
            ["createdAt", "Date", "Default: Now", "Flashcard creation timestamp"]
        ]),
        ("Table 9: Formula Schema (formulas)", [
            ["Field Name", "Data Type", "Constraint", "Description"],
            ["_id", "ObjectId", "Primary Key", "Unique formula entry identifier"],
            ["id", "String", "Unique, Required", "Formula reference code"],
            ["category", "String", "Required", "Topic (Profit & Loss, Algebra, Calculus)"],
            ["title", "String", "Required", "Cheatsheet title (English & Gujarati)"],
            ["grades", "Array[String]", "Optional", "Applicable student grades"],
            ["formulas", "Array[String]", "Required", "Mathematical equations & expressions"],
            ["notes", "String", "Optional", "Important conditions and memory tips"],
            ["createdAt", "Date", "Default: Now", "Reference creation timestamp"]
        ]),
        ("Table 10: ImportantQuestion Schema (importantquestions)", [
            ["Field Name", "Data Type", "Constraint", "Description"],
            ["_id", "ObjectId", "Primary Key", "Unique question bank identifier"],
            ["id", "String", "Unique, Required", "Reference question identifier"],
            ["grade", "String", "Required, Indexed", "Target grade level (Class 9-12, college)"],
            ["board", "String", "Default: 'All'", "Board designation (GSEB, CBSE)"],
            ["subject", "String", "Required, Indexed", "Academic or placement subject"],
            ["chapterNumber", "Number", "Default: 1", "Chapter number in curriculum"],
            ["chapterTitle", "String", "Required", "Full chapter title"],
            ["marks", "Number", "Required, Indexed", "Question weight (1, 2, 3, 4, 5 marks)"],
            ["marksCategory", "String", "Required", "Category (VSA, Short, Long Answer)"],
            ["questionText", "String", "Required", "Official examination question"],
            ["modelAnswer", "String", "Required", "Complete official step-by-step solution"]
        ]),
        ("Table 11: Curriculum Schema (curriculums)", [
            ["Field Name", "Data Type", "Constraint", "Description"],
            ["_id", "ObjectId", "Primary Key", "Unique curriculum identifier"],
            ["grade", "String", "Unique, Required", "Academic grade identifier ('10', 'college')"],
            ["title", "String", "Required", "Full curriculum track title"],
            ["description", "String", "Optional", "Overview of educational syllabus"],
            ["badge", "String", "Optional", "Display emblem / badge name"],
            ["subjects", "Array[Object]", "Required", "Subjects with nested chapters & videos"],
            ["createdAt", "Date", "Default: Now", "Curriculum release timestamp"]
        ])
    ]

    # Find "3.5 Data Dictionary" and "4. Development"
    p_dd = None
    p_dev = None
    for p in doc.paragraphs:
        if p.text.strip() == "3.5 Data Dictionary":
            p_dd = p
        elif p.text.strip() == "4. Development":
            p_dev = p
            break

    assert p_dd is not None and p_dev is not None, "Could not locate 3.5 or 4.0 paragraphs"

    # Remove all existing tables between 3.5 and 4.0
    # Also remove any paragraphs between 3.5 and 4.0 (like old Table 1..4 titles)
    # Let's collect elements between p_dd and p_dev
    curr = p_dd._element.getnext()
    to_delete = []
    while curr is not None and curr != p_dev._element:
        nxt = curr.getnext()
        to_delete.append(curr)
        curr = nxt

    print(f"Removing {len(to_delete)} old elements between 3.5 and 4.0...")
    for elem in to_delete:
        elem.getparent().remove(elem)

    # Now sequentially insert all 11 schemas before p_dev
    for title_text, rows_data in all_schemas:
        # Title paragraph
        p_t = p_dev.insert_paragraph_before()
        p_t.paragraph_format.space_before = Pt(14)
        p_t.paragraph_format.space_after = Pt(4)
        p_t.paragraph_format.keep_with_next = True
        r = p_t.add_run(title_text)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(30, 58, 138)

        # Table
        tbl = doc.add_table(rows=len(rows_data), cols=4)
        populate_table(tbl, rows_data)
        # Move tbl before p_dev
        p_dev._element.addprevious(tbl._element)

        # Spacing paragraph after table
        p_space = p_dev.insert_paragraph_before()
        p_space.paragraph_format.space_before = Pt(2)
        p_space.paragraph_format.space_after = Pt(6)

    doc.save(doc_path)
    print(f"Successfully updated {doc_path} with all 11 collections! Total tables in doc: {len(doc.tables)}")

if __name__ == "__main__":
    update_data_dictionary()

