import openpyxl

wb = openpyxl.load_workbook('Agile Document (Recovered).xlsx', data_only=True)
for name in wb.sheetnames:
    ws = wb[name]
    print(f"\n==================== SHEET: {name} ====================")
    for r in range(1, ws.max_row + 1):
        row = [str(ws.cell(r, c).value or '').strip() for c in range(1, ws.max_column + 1)]
        # remove trailing empty strings
        while row and not row[-1]:
            row.pop()
        if row:
            print(f"R{r:2d}: " + " | ".join(row))

