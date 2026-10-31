import pptx
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
import shutil

all_schemas = [
    ("Table 1: User Schema (users)", [
        ["Field Name", "Data Type", "Constraint", "Description"],
        ["_id", "ObjectId", "Primary Key", "Unique user identifier"],
        ["name", "String", "Required", "Full name of the student"],
        ["email", "String", "Unique, Required", "Student email address"],
        ["passwordHash", "String", "Required", "Bcrypt-hashed password string"],
        ["activeTrack", "String", "Default: 'placement'", "Selected track ('placement', 'school')"],
        ["grade", "String", "Enum: 6-12, 'college'", "Current academic grade / year"],
        ["board", "String", "Enum: CBSE, GSEB, ICSE", "Academic education board"],
        ["medium", "String", "Enum: English, Gujarati", "Medium of instruction"],
        ["streakCount", "Number", "Default: 3", "Daily consecutive study streak"],
        ["targetCompany", "String", "Default: 'All Companies'", "Target placement employer"],
        ["createdAt", "Date", "Default: Now", "Account creation timestamp"]
    ]),
    ("Table 2: Question Schema (questions)", [
        ["Field Name", "Data Type", "Constraint", "Description"],
        ["_id", "ObjectId", "Primary Key", "Unique question identifier"],
        ["id", "Number", "Unique, Indexed", "Numeric question index"],
        ["questionText", "String", "Required", "The practice question statement"],
        ["options", "Array[String]", "Required, Len: 4", "Multiple-choice answer options"],
        ["correctAnswer", "Number", "Required", "0-based index of correct option"],
        ["category", "String", "Required, Indexed", "Aptitude, Logical, Technical topic"],
        ["company", "String", "Optional, Indexed", "Target company tag (TCS, Infosys)"],
        ["difficulty", "String", "Enum: Easy, Med, Hard", "Calibrated difficulty level"],
        ["explanation", "String", "Optional", "Detailed step-by-step solution"],
        ["grades", "Array[String]", "Indexed", "Target grades (['college'], ['10'])"]
    ]),
    ("Table 3: Course Schema (courses)", [
        ["Field Name", "Data Type", "Constraint", "Description"],
        ["_id", "ObjectId", "Primary Key", "Unique course identifier"],
        ["title", "String", "Required", "Title of the IT course / subject"],
        ["category", "String", "Required", "Course domain (DSA, Web Dev, Cloud)"],
        ["instructor", "String", "Required", "Course instructor / educator name"],
        ["duration", "String", "Default: '10 hrs'", "Total instructional video hours"],
        ["lessonsCount", "Number", "Default: 20", "Total number of video lessons"],
        ["rating", "Number", "Default: 4.8", "Student rating out of 5.0"],
        ["tag", "String", "Default: 'Trending'", "Display badge tag"],
        ["level", "String", "Default: 'All Levels'", "Target student skill level"],
        ["thumbnailUrl", "String", "Optional", "Course cover graphic image URL"]
    ]),
    ("Table 4: ExamResult Schema (examresults)", [
        ["Field Name", "Data Type", "Constraint", "Description"],
        ["_id", "ObjectId", "Primary Key", "Unique result identifier"],
        ["userId", "ObjectId", "Foreign Key (User)", "Reference to student account"],
        ["userEmail", "String", "Default: Guest", "Student email address"],
        ["grade", "String", "Default: 'college'", "Academic grade level"],
        ["board", "String", "Default: 'CBSE'", "Examination board context"],
        ["category", "String", "Default: 'General'", "Subject or company test category"],
        ["scorePercentage", "Number", "Required", "Calculated score percentage (0-100)"],
        ["correctCount", "Number", "Required", "Number of correctly answered questions"],
        ["totalQuestions", "Number", "Required", "Total test questions count"],
        ["userAnswers", "Object", "Default: {}", "Key-value pairs of student answers"],
        ["weakCategories", "Array[String]", "Optional", "Topics where accuracy fell <60%"],
        ["createdAt", "Date", "Default: Now", "Exam submission timestamp"]
    ]),
    ("Table 5: DuelRoom Schema (duelrooms)", [
        ["Field Name", "Data Type", "Constraint", "Description"],
        ["_id", "ObjectId", "Primary Key", "Unique duel room identifier"],
        ["roomCode", "String", "Unique, Indexed", "6-character room code for 1v1 PvP"],
        ["grade", "String", "Required", "Academic grade level"],
        ["subject", "String", "Required, Default: 'All'", "Test subject filter"],
        ["status", "String", "Enum, Indexed", "waiting, ready, active, completed, abandoned"],
        ["host", "Object", "Required", "Host details (name, avatar, score, ready)"],
        ["guest", "Object", "Default: null", "Opponent student details"],
        ["questions", "Array[Object]", "Required", "Synchronized test questions for duel"],
        ["createdAt", "Date", "TTL: 3600s (1 hr)", "Auto-expiring session timestamp"]
    ]),
    ("Table 6: SamplePaper Schema (samplepapers)", [
        ["Field Name", "Data Type", "Constraint", "Description"],
        ["_id", "ObjectId", "Primary Key", "Unique sample paper identifier"],
        ["id", "String", "Unique, Required", "Model paper identifier code"],
        ["title", "String", "Required", "Official model paper title"],
        ["grade", "String", "Required, Indexed", "Target grade ('10', '12', 'college')"],
        ["board", "String", "Default: 'CBSE', Indexed", "Board (CBSE, GSEB, Placement)"],
        ["subject", "String", "Required, Indexed", "Examination subject"],
        ["year", "String", "Default: '2024-2025'", "Academic examination year"],
        ["durationMinutes", "Number", "Default: 60", "Total allocated exam duration"],
        ["totalMarks", "Number", "Default: 50", "Total examination marks"],
        ["questionCount", "Number", "Default: 25", "Total question count"],
        ["officialPdfUrl", "String", "Optional", "URL to downloadable paper PDF"]
    ]),
    ("Table 7: VideoLecture Schema (videolectures)", [
        ["Field Name", "Data Type", "Constraint", "Description"],
        ["_id", "ObjectId", "Primary Key", "Unique lecture identifier"],
        ["id", "String", "Unique, Required", "Video lesson identifier code"],
        ["grades", "Array[String]", "Indexed", "Grade tags (['college'], ['10'])"],
        ["boards", "Array[String]", "Indexed", "Educational boards (CBSE, GSEB)"],
        ["category", "String", "Required, Indexed", "Subject or technical topic category"],
        ["title", "String", "Required", "Lecture lesson title"],
        ["instructor", "String", "Required", "Course educator name"],
        ["duration", "String", "Required", "Formatted length (e.g. '18m 30s')"],
        ["durationSeconds", "Number", "Required", "Total playback duration in seconds"],
        ["youtubeVideoId", "String", "Required", "Embedded YouTube video ID"],
        ["thumbnailUrl", "String", "Required", "Thumbnail image URL"]
    ]),
    ("Table 8: Flashcard Schema (flashcards)", [
        ["Field Name", "Data Type", "Constraint", "Description"],
        ["_id", "ObjectId", "Primary Key", "Unique flashcard identifier"],
        ["id", "String", "Unique, Required", "Concept flashcard code"],
        ["topic", "String", "Required", "Subject topic (SQL, OS, Aptitude)"],
        ["question", "String", "Required", "Concept prompt or interview question"],
        ["answer", "String", "Required", "Concise verified answer / definition"],
        ["interviewerTips", "String", "Optional", "Practical advice for campus interview"],
        ["difficulty", "String", "Enum: Beginner, Med, Adv", "Topic complexity level"],
        ["tags", "Array[String]", "Optional", "Search keywords and category tags"],
        ["createdAt", "Date", "Default: Now", "Flashcard creation timestamp"]
    ]),
    ("Table 9: Formula Schema (formulas)", [
        ["Field Name", "Data Type", "Constraint", "Description"],
        ["_id", "ObjectId", "Primary Key", "Unique formula entry identifier"],
        ["id", "String", "Unique, Required", "Formula reference code"],
        ["category", "String", "Required", "Topic (Profit & Loss, Algebra, Calculus)"],
        ["title", "String", "Required", "Cheatsheet title (English & Gujarati)"],
        ["grades", "Array[String]", "Optional", "Applicable student grades"],
        ["formulas", "Array[String]", "Required", "Mathematical equations & expressions"],
        ["notes", "String", "Optional", "Important conditions and memory tips"],
        ["createdAt", "Date", "Default: Now", "Reference creation timestamp"]
    ]),
    ("Table 10: ImportantQuestion Schema (importantquestions)", [
        ["Field Name", "Data Type", "Constraint", "Description"],
        ["_id", "ObjectId", "Primary Key", "Unique question bank identifier"],
        ["id", "String", "Unique, Required", "Reference question identifier"],
        ["grade", "String", "Required, Indexed", "Target grade level (Class 9-12, college)"],
        ["board", "String", "Default: 'All'", "Board designation (GSEB, CBSE)"],
        ["subject", "String", "Required, Indexed", "Academic or placement subject"],
        ["chapterNumber", "Number", "Default: 1", "Chapter number in curriculum"],
        ["chapterTitle", "String", "Required", "Full chapter title"],
        ["marks", "Number", "Required, Indexed", "Question weight (1, 2, 3, 4, 5 marks)"],
        ["marksCategory", "String", "Required", "Category (VSA, Short, Long Answer)"],
        ["questionText", "String", "Required", "Official examination question"],
        ["modelAnswer", "String", "Required", "Complete official step-by-step solution"]
    ]),
    ("Table 11: Curriculum Schema (curriculums)", [
        ["Field Name", "Data Type", "Constraint", "Description"],
        ["_id", "ObjectId", "Primary Key", "Unique curriculum identifier"],
        ["grade", "String", "Unique, Required", "Academic grade identifier ('10', 'college')"],
        ["title", "String", "Required", "Full curriculum track title"],
        ["description", "String", "Optional", "Overview of educational syllabus"],
        ["badge", "String", "Optional", "Display emblem / badge name"],
        ["subjects", "Array[Object]", "Required", "Subjects with nested chapters & videos"],
        ["createdAt", "Date", "Default: Now", "Curriculum release timestamp"]
    ])
]

def format_theme_table(slide, subtitle_text, rows_data):
    # Remove all non-title shapes
    for shp in [s for s in slide.shapes if s != slide.shapes.title]:
        sp = shp._element
        sp.getparent().remove(sp)

    # Set Title
    if slide.shapes.title:
        slide.shapes.title.text = "3.5 Data Dictionary"
        slide.shapes.title.left = Inches(0.59)
        slide.shapes.title.top = Inches(0.55)
        slide.shapes.title.width = Inches(8.69)
        slide.shapes.title.height = Inches(0.80)

    # Add Subtitle TextBox (like original Slide 20)
    tb = slide.shapes.add_textbox(Inches(0.65), Inches(1.55), Inches(8.7), Inches(0.45))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = subtitle_text
    p.font.bold = True
    p.font.size = Pt(13)

    # Add Table with native theme table style {F2DE63D5-997A-4646-A377-4702673A728D}
    tbl_left = Inches(1.15)
    tbl_top = Inches(2.15)
    tbl_width = Inches(7.70)
    tbl_height = Inches(min(3.2, 0.28 * len(rows_data) + 0.3))

    tbl_shape = slide.shapes.add_table(len(rows_data), 4, tbl_left, tbl_top, tbl_width, tbl_height)
    tbl = tbl_shape.table
    
    # Configure table properties to match previous table exactly
    tblPr = tbl._tbl.tblPr
    tblPr.set('firstRow', '1')
    tblPr.set('bandRow', '1')
    tblPr.set('firstCol', '1')
    
    style_node = tblPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}tableStyleId')
    if style_node is not None:
        style_node.text = '{F2DE63D5-997A-4646-A377-4702673A728D}'

    col_widths = [Inches(1.65), Inches(1.45), Inches(1.70), Inches(2.90)]
    for c_idx in range(4):
        tbl.columns[c_idx].width = col_widths[c_idx]

    # Populate cell text WITHOUT overriding cell.fill (so theme table style applies natively!)
    for r_idx, row_values in enumerate(rows_data):
        for c_idx, val in enumerate(row_values):
            cell = tbl.cell(r_idx, c_idx)
            cell.text = val
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)
            cell.margin_left = Inches(0.08)
            cell.margin_right = Inches(0.08)

            p = cell.text_frame.paragraphs[0]
            p.word_wrap = True
            if r_idx == 0:
                p.font.bold = True
                p.font.size = Pt(9.5)
            else:
                p.font.bold = False
                p.font.size = Pt(8.5)

def update_presentation_tables():
    pptx_path = "Review_Presentation.pptx"
    prs = pptx.Presentation(pptx_path)
    
    print(f"Updating 11 schema slides in {pptx_path} (Slides 20 to 30)...")
    for i in range(11):
        slide = prs.slides[19 + i]
        subtitle, rows = all_schemas[i]
        format_theme_table(slide, subtitle, rows)
        print(f"  Slide {20 + i}: {subtitle} updated with native theme table style!")

    prs.save(pptx_path)
    print("Successfully updated Review_Presentation.pptx with previous table style!")

if __name__ == "__main__":
    update_presentation_tables()

