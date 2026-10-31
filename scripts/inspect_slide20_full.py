import pptx

prs = pptx.Presentation('Review_Presentation_backup_orig.pptx')
s20 = prs.slides[19]
print("Number of shapes in slide 20:", len(s20.shapes))
for idx, shp in enumerate(s20.shapes):
    print(f"\nShape {idx}: name='{shp.name}', type={shp.shape_type}")
    print(f"  left={shp.left.inches:.3f}, top={shp.top.inches:.3f}, width={shp.width.inches:.3f}, height={shp.height.inches:.3f}")
    if shp.has_text_frame:
        print(f"  Text: {repr(shp.text_frame.text)}")
        for p in shp.text_frame.paragraphs:
            print(f"    P: font={p.font.name}, size={p.font.size.pt if p.font.size else None}, bold={p.font.bold}, color={p.font.color.rgb if p.font.color and p.font.color.type else None}, align={p.alignment}")
    if shp.has_table:
        t = shp.table
        print(f"  Table: {len(t.rows)} rows, {len(t.columns)} cols")
        for c_idx in range(len(t.columns)):
            print(f"    Col {c_idx} width: {t.columns[c_idx].width.inches:.3f} in")
        for r_idx in range(len(t.rows)):
            row_txt = [c.text.strip() for c in t.rows[r_idx].cells]
            print(f"    Row {r_idx}: " + " | ".join(row_txt))

