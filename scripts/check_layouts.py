import pptx

prs = pptx.Presentation('Review_Presentation_backup_orig.pptx')
for idx, s in enumerate(prs.slides):
    t = s.shapes.title.text.strip().replace('\n', ' ') if s.shapes.title else ''
    print(f"Slide {idx+1:2d} (idx {idx:2d}): layout {s.slide_layout.name} - '{t}'")

