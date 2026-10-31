import docx

doc = docx.Document('ProjectReviewSheet-ReactNative.docx')
t = doc.tables[0]

print("R2 cell 0 text:", repr(t.cell(2, 0).text))
print("R2 cell 1 text:", repr(t.cell(2, 1).text))
print("R8 cell 0 text:", repr(t.cell(8, 0).text))
print("R8 cell 3 text:", repr(t.cell(8, 3).text))
print("R9 cell 3 text:", repr(t.cell(9, 3).text))
print("R10 cell 3 text:", repr(t.cell(10, 3).text))

