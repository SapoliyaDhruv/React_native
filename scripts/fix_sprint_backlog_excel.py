import openpyxl

wb = openpyxl.load_workbook('Agile Document (Recovered).xlsx')
ws = wb['Product Backlog']

# Let's clean up rows 18 to 47 to match MindSprint sprints accurately
sprint_backlog_data = [
    # Task Name, Story, Sprint Ready, Priority, Status, Story Point, Assign to sprint
    ("Sprint 1: Architecture, Navigation & Auth", "Yes", "Yes", "High", "Completed", 16, "Dhruv"),
    ("Sprint Planning & Stack Setup", "Yes", "Yes", "High", "Completed", 3, "Dhruv"),
    ("Login & Registration Backend (JWT)", "Yes", "Yes", "High", "Completed", 5, "Dvarkesh"),
    ("Authentication UI & Navigation Routing", "Yes", "Yes", "High", "Completed", 8, "Krunal"),

    ("Sprint 2: Course Catalog & Video Masterclasses", "Yes", "Yes", "High", "Completed", 21, "Dhruv"),
    ("Syllabus Selector & User Profile Module", "Yes", "Yes", "High", "Completed", 3, "Dhruv"),
    ("Course & Video Catalog REST APIs", "Yes", "Yes", "High", "Completed", 8, "Dhruv,Dvarkesh"),
    ("Course Catalog UI & Chapter Reader", "Yes", "Yes", "High", "Completed", 10, "Krunal"),

    ("Sprint 3: Timed Mock Exam Engine & Palette", "Yes", "Yes", "High", "Completed", 26, "Dhruv"),
    ("Question Bank Backend & Company Tagging", "Yes", "Yes", "High", "Completed", 8, "Krunal,Dvarkesh"),
    ("Timed Mock Exam Engine UI & Palette", "Yes", "Yes", "High", "Completed", 13, "Dhruv,Dvarkesh"),
    ("In-Exam Scratchpad & Timer Logic", "Yes", "Yes", "High", "Completed", 5, "Dhruv"),

    ("Sprint 4: Weak Spot Detection & Result Evaluation", "Yes", "Yes", "High", "Completed", 21, "Dhruv"),
    ("Score Evaluation & Percentage Logic", "Yes", "Yes", "High", "Completed", 5, "Krunal"),
    ("Weak Spot Detection & Accuracy Analytics", "Yes", "Yes", "High", "Completed", 8, "Dhruv,Krunal"),
    ("Results Summary UI & Category Cards", "Yes", "Yes", "High", "Completed", 8, "Dvarkesh,Krunal"),

    ("Sprint 5: AI Doubt Solver & Flashcards", "Yes", "Yes", "High", "Completed", 24, "Dhruv"),
    ("AI Doubt Solver Text & Voice Processing", "Yes", "Yes", "High", "Completed", 8, "Dvarkesh"),
    ("Camera OCR Text Extraction & Matching", "Yes", "Yes", "High", "Completed", 8, "Dhruv,Krunal"),
    ("Swipeable Flashcards & Daily Streak", "Yes", "Yes", "High", "Completed", 8, "Krunal,Dvarkesh"),

    ("Sprint 6: Offline Resilience & UAT", "Yes", "Yes", "High", "Completed", 20, "Dhruv"),
    ("Offline Resilience & Local Fallback Layer", "Yes", "Yes", "High", "Completed", 8, "Krunal"),
    ("UAT Testing, Bug Fixes & Optimization", "Yes", "Yes", "High", "Completed", 5, "Dhruv,Dvarkesh"),
    ("Final Project Documentation & Review Prep", "Yes", "Yes", "High", "Completed", 7, "Dhruv,Krunal,Dvarkesh")
]

start_row = 18
for i, item in enumerate(sprint_backlog_data):
    curr_r = start_row + i
    ws.cell(curr_r, 2).value = item[0]
    ws.cell(curr_r, 3).value = item[1]
    ws.cell(curr_r, 4).value = item[2]
    ws.cell(curr_r, 5).value = item[3]
    ws.cell(curr_r, 6).value = item[4]
    ws.cell(curr_r, 7).value = item[5]
    ws.cell(curr_r, 8).value = item[6]

# Clear any extra rows below
for r in range(start_row + len(sprint_backlog_data), ws.max_row + 1):
    for c in range(1, ws.max_column + 1):
        ws.cell(r, c).value = None

wb.save('Agile Document (Recovered).xlsx')
print('Successfully aligned Sprint Backlog in Excel with MindSprint!')

