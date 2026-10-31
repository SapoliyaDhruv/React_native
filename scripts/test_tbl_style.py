import pptx
from pptx.util import Inches

prs = pptx.Presentation('Review_Presentation_backup_orig.pptx')
s = prs.slides.add_slide(prs.slide_layouts[1])
tbl_shape = s.shapes.add_table(3, 4, Inches(1.25), Inches(2.2), Inches(7.5), Inches(2.0))
tbl = tbl_shape.table
tbl._tbl.tblPr.set('firstCol', '1')
style_node = tbl._tbl.tblPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}tableStyleId')
if style_node is not None:
    style_node.text = '{F2DE63D5-997A-4646-A377-4702673A728D}'

print("Updated tblPr xml:\n", tbl._tbl.tblPr.xml)
