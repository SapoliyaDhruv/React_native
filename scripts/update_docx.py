import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

def update_docx():
    doc_path = "MindSprint_Project_Documentation.docx"
    doc = docx.Document(doc_path)
    print(f"Loaded {doc_path} with {len(doc.paragraphs)} paragraphs")

    # 1. Update images in Section 5 (Agile Documentation)
    # Paragraph mapping for agile images:
    # P173: 5.1 Project Charter
    # P177: 5.2 Project Roadmap
    # P181: 5.3 Project Plan
    # P185: 5.4 User Stories
    # P189: 5.5 Release Plan
    # P193: 5.6 Product Backlog
    # P197: 5.7 Test Plan
    # P201: 5.8 Burndown Chart
    agile_image_map = [
        (173, "screenshots/agile_01_project_charter.png"),
        (177, "screenshots/agile_02_project_roadmap.png"),
        (181, "screenshots/agile_03_project_plan.png"),
        (185, "screenshots/agile_04_user_stories.png"),
        (189, "screenshots/agile_05_release_plan.png"),
        (193, "screenshots/agile_06_product_backlog.png"),
        (197, "screenshots/agile_07_test_plan.png"),
        (201, "screenshots/agile_08_burndown_chart.png")
    ]

    for p_idx, img_path in agile_image_map:
        p = doc.paragraphs[p_idx]
        p.text = "" # Clear old runs and images
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(img_path, width=Inches(6.2))
        print(f"Replaced image at paragraph {p_idx} with {img_path}")

    # Helper function for adding styled headings and text
    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(30, 58, 138)
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(15, 23, 42)
        return p

    def add_body(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(51, 65, 85)
        return p

    def add_bullet(bold_prefix, text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        run1 = p.add_run(bold_prefix + " ")
        run1.bold = True
        run1.font.name = "Calibri"
        run1.font.size = Pt(11)
        run1.font.color.rgb = RGBColor(30, 58, 138)
        run2 = p.add_run(text)
        run2.font.name = "Calibri"
        run2.font.size = Pt(11)
        run2.font.color.rgb = RGBColor(51, 65, 85)
        return p

    # ========================================================
    # 6. Proposed Enhancements
    # ========================================================
    add_h1("6. Proposed Enhancements")
    add_body(
        "While the current release of MindSprint delivers an end-to-end placement preparation, "
        "timed test simulation, and video learning ecosystem, several high-value architectural "
        "and functional enhancements are planned for subsequent iterations based on feedback gathered "
        "during user acceptance testing:"
    )

    add_h2("6.1 Adaptive AI Difficulty Engine & Personalized Study Paths")
    add_body(
        "Currently, mock exams draw questions based on static difficulty levels (Easy, Medium, Hard). "
        "The proposed enhancement incorporates an Item Response Theory (IRT) adaptive testing engine. "
        "As a student answers questions during an active test session, the system will dynamically adjust "
        "subsequent question difficulty in real time based on their response accuracy, elapsed seconds per "
        "question, and category weakness history. This will provide placement aspirants with truly calibrated "
        "mock exams mirroring adaptive recruitment tests conducted by tier-1 tech firms."
    )

    add_h2("6.2 Real-Time Multiplayer Quiz Battles & Collaborative Study Rooms")
    add_body(
        "The existing 1v1 Quiz Duel module allows paired peer battles through room codes. The next release "
        "will expand this into multi-user live classroom tournaments supporting up to 50 concurrent students. "
        "Additionally, collaborative virtual study rooms with integrated WebRTC audio channels will enable peer "
        "groups to solve complex technical problems and practice mock technical viva interviews together in real time."
    )

    add_h2("6.3 Institutional Placement Cell & Faculty Analytics Web Dashboard")
    add_body(
        "To maximize institutional utility across colleges and training institutes, a dedicated React/Next.js "
        "administrative web portal will be developed for college placement officers and subject professors. "
        "This portal will offer batch-level analytics, cohort readiness tracking, automated weak topic heatmaps "
        "across departments, custom test scheduling, and automated PDF export of student placement readiness reports."
    )

    add_h2("6.4 Background Test Sync & Smart Push Notification Reminders")
    add_body(
        "The current offline resilience architecture gracefully serves local questions and allows offline exam taking. "
        "The proposed enhancement will implement an automatic background synchronization worker using React Native "
        "Background Actions. When a student completes an offline test during a commute and connectivity is subsequently "
        "detected, the test results, flagged bookmarks, and streak milestones will sync silently to MongoDB without "
        "requiring user interaction. Smart push notifications will also alert students to maintain daily revision streaks."
    )

    add_h2("6.5 Regional Vernacular Language Support (Gujarati & Hindi)")
    add_body(
        "To improve accessibility for Gujarati medium school students (GSEB Class 9-12) and regional engineering diploma "
        "students, bilingual question explanations, audio viva transcripts, and formula walkthroughs will be introduced "
        "in Gujarati and Hindi alongside standard English terminology."
    )

    # ========================================================
    # 7. Conclusion
    # ========================================================
    add_h1("7. Conclusion")

    add_h2("7.1 Summary of Project Achievements")
    add_body(
        "MindSprint was conceived and built to eliminate the fragmentation, high cost, and lack of personalization "
        "in traditional student placement preparation. Over the course of six structured Agile Scrum sprints, the team "
        "successfully designed, engineered, tested, and documented a complete full-stack mobile learning ecosystem."
    )
    add_body(
        "The application integrates real-time timed mock tests with full question palettes, automated scoring algorithms "
        "that instantly diagnose student weaknesses (<60% accuracy threshold), curated video masterclasses with chapter "
        "timings, official Gujarat Board (GSEB) and GTU engineering digital textbooks, swipeable spaced repetition flashcards, "
        "and an AI-powered Doubt Solver with camera OCR capabilities."
    )

    add_h2("7.2 Technical Realization & Engineering Impact")
    add_body(
        "From an engineering standpoint, the project demonstrates robust architectural design patterns:"
    )
    add_bullet("Cross-Platform Portability:", "Developed on React Native and Expo (SDK v54/v57) with TypeScript, ensuring high performance, responsive layout, and identical user experiences across Android mobile devices, iOS, and Web browsers.")
    add_bullet("Full-Stack Decoupling:", "A clean separation of concerns between the mobile client and the Node.js/Express REST backend, secured with bcrypt password hashing and stateless JWT bearer authentication.")
    add_bullet("Zero-Downtime Offline Resilience:", "The ApiService fallback layer guarantees that students never encounter broken screens or crashes when studying in low-connectivity environments, transparently loading verified static datasets.")
    add_bullet("Agile Scrum Rigor:", "Iterative delivery validated through user stories, milestone release increments, continuous backlog grooming, and comprehensive regression test plans.")

    add_h2("7.3 Concluding Remarks")
    add_body(
        "In conclusion, MindSprint provides a comprehensive, student-centric digital companion that directly empowers "
        "college placement aspirants and school students across Gujarat to enhance their technical problem-solving capabilities, "
        "build disciplined study habits, and achieve competitive success in campus recruitment drives and board examinations."
    )

    # ========================================================
    # 8. Bibliography
    # ========================================================
    add_h1("8. Bibliography")
    add_body(
        "The development of MindSprint referenced official technical specifications, academic curriculum frameworks, "
        "software engineering literature, and developer documentation as detailed below:"
    )

    references = [
        ("1. Expo Framework Documentation:", "Expo Development Team. (2026). Expo SDK Documentation, Architecture Guides, and Expo Router Reference. Available at: https://docs.expo.dev"),
        ("2. React Native Documentation:", "Meta Platforms, Inc. (2026). React Native: Learn once, write anywhere. Core Components, Hooks API, and New Architecture. Available at: https://reactnative.dev"),
        ("3. Node.js & Express.js Developer Guides:", "OpenJS Foundation. (2026). Node.js v20 LTS Runtime Environment & Express.js RESTful API Framework Reference Manual. Available at: https://expressjs.com"),
        ("4. MongoDB & Mongoose Object Data Modeling:", "MongoDB, Inc. (2026). Mongoose ODM v8 Documentation: Schema Definitions, Validation, Indexing, and Query Middleware. Available at: https://mongoosejs.com"),
        ("5. Google Gemini AI API Documentation:", "Google DeepMind. (2026). Gemini API Multimodal Developer Reference: Vision Question Parsing and Generative AI SDK. Available at: https://ai.google.dev"),
        ("6. Agile Scrum Framework Guide:", "Schwaber, K., & Sutherland, J. (2020). The Scrum Guide: The Definitive Guide to Scrum: The Rules of the Game. Scrum.org."),
        ("7. Gujarat Technological University (GTU) Academic Syllabus:", "GTU Academic Council. (2025). Bachelor of Engineering Curriculum Guidelines, Aptitude & Technical Skill Assessment Frameworks. Ahmedabad, Gujarat."),
        ("8. Gujarat Secondary and Higher Secondary Education Board (GSEB):", "GSEB Gandhinagar. (2025). State Board Syllabi, Blueprints, and Official Reference Textbooks for Mathematics and Science.")
    ]

    for prefix, citation in references:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.line_spacing = 1.15
        r1 = p.add_run(prefix + " ")
        r1.bold = True
        r1.font.name = "Calibri"
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = RGBColor(30, 58, 138)
        r2 = p.add_run(citation)
        r2.font.name = "Calibri"
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = RGBColor(51, 65, 85)

    doc.save("MindSprint_Project_Documentation.docx")
    print(f"MindSprint_Project_Documentation.docx updated successfully! Total paragraphs: {len(doc.paragraphs)}")

if __name__ == "__main__":
    update_docx()

